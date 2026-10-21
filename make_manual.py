# -*- coding: utf-8 -*-
"""生成软件说明书（Word，含界面截图与操作步骤）。"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = r'D:\桌面\APP'
SHOTS = os.path.join(ROOT, '交付物', 'screenshots')
OUT = os.path.join(ROOT, '交付物', '软件说明书.docx')

doc = Document()

# 页面设置：A4
for section in doc.sections:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

# 默认字体
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(12)
normal._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')


def zh(run, font='宋体'):
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font)


def para(text, size=12, bold=False, align=None, font='宋体', space_after=6, indent=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    zh(r, font)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.5
    if indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    return p


def h1(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(16)
    r.bold = True
    r.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
    zh(r, '微软雅黑')
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    return p


def h2(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(14)
    r.bold = True
    zh(r, '微软雅黑')
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    return p


def img(path, caption, width_cm=15.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(os.path.join(SHOTS, path), width=Cm(width_cm))
    para(caption, size=10, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, font='微软雅黑')


def page_break():
    doc.add_page_break()


# ===================== 封面 =====================
for _ in range(4):
    doc.add_paragraph()
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('高德路径搜索与下载工具'); r.font.size = Pt(28); r.bold = True; zh(r, '微软雅黑')
t2 = doc.add_paragraph(); t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t2.add_run('软 件 说 明 书'); r.font.size = Pt(22); r.bold = True; zh(r, '微软雅黑')
for _ in range(6):
    doc.add_paragraph()
para('版本：1.0', size=14, align=WD_ALIGN_PARAGRAPH.CENTER, font='微软雅黑')
para('适用平台：Windows（本地运行）', size=14, align=WD_ALIGN_PARAGRAPH.CENTER, font='微软雅黑')
page_break()

# ===================== 一、软件简介 =====================
h1('一、软件简介')
para('高德路径搜索与下载工具是一款在本地浏览器中运行的轻量级网页应用，基于高德开放平台的路径规划接口，'
     '帮助用户搜索驾车、步行、骑行、公交等出行方式的路线，并在交互式地图上直观展示。', indent=True)
para('本软件的核心价值在于：不仅能够"搜路线"，还支持把路线数据导出为通用的 CSV 与 GeoJSON 格式（统一转换为 WGS84 坐标系），'
     '便于在其它 GIS 软件、GPS 设备或数据分析工具中继续使用。同时提供坐标互转、Excel 批量导入、里程统计、今日计划、'
     '轨迹录制与离线瓦片缓存等实用功能。', indent=True)
para('主要功能如下：', indent=False)
for f in [
    '① 路线搜索：支持驾车、步行、骑行、公交四种方式，可输入地址或经纬度（默认 WGS84）。',
    '② 驾车策略：最快捷、最短距离、少收费、躲避拥堵、不走高速等可选。',
    '③ 途经点：驾车支持在起终点之间插入多个途经点。',
    '④ 实时路况：驾车路线按畅通、缓行、拥堵、严重拥堵着色显示。',
    '⑤ 结果展示：总距离、预计耗时、分段导航指令表。',
    '⑥ 导出：单条路线 CSV、GeoJSON（WGS84）。',
    '⑦ 坐标互转：GCJ-02 ↔ WGS84。',
    '⑧ 轨迹录制：浏览器定位采样，导出轨迹 GeoJSON / CSV。',
    '⑨ 批量导入：Excel / CSV 一次搜索多条、一次下载多条、里程统计表。',
    '⑩ 今日计划：当日路线任务管理、排序、一键导出。',
    '⑪ 备份恢复：一键打包 data/ 与 downloads/ 为 zip，或从 zip 恢复。',
    '⑫ 离线地图：浏览过的地图瓦片自动缓存。',
]:
    para(f, size=12, indent=True, space_after=3)
page_break()

# ===================== 二、运行环境 =====================
h1('二、运行环境')
h2('2.1 硬件环境')
para('普通台式机或笔记本电脑即可，无特殊硬件要求。')
h2('2.2 软件环境')
for f in [
    '操作系统：Windows 10 / 11（本说明书基于 Windows 11 编写）。',
    'Python：3.8 及以上版本（需勾选 "Add Python to PATH"）。',
    '依赖库：openpyxl（用于 Excel 批量导入；不使用 Excel 时可不安装）。',
    '浏览器：Edge、Chrome 等现代浏览器。',
    '网络：仅"搜索路线""公交/逆地理编码""加载地图底图"需要联网。',
]:
    para(f, size=12, space_after=3)
page_break()

# ===================== 三、安装与启动 =====================
h1('三、安装与启动')
h2('3.1 获取源码')
para('将源码压缩包解压到任意目录（本说明书以 D:\\桌面\\APP 为例）。源码包含：'
     'app.py（后端主程序）、amap_client.py（高德接口封装）、coord.py（坐标互转）、'
     'storage.py（本地持久化）、run_tests.py（测试运行器）、static/（前端页面）、'
     'tests/（测试）、requirements.txt、README.md 等。', indent=True)
h2('3.2 安装依赖')
para('打开命令行，进入项目目录，执行：', indent=True)
para('pip install -r requirements.txt', size=12, font='微软雅黑', indent=True)
para('若不需要 Excel 批量导入，也可以不安装 openpyxl，改用 .csv 文件导入。', indent=True)
h2('3.3 启动')
para('在项目目录执行：', indent=True)
para('python app.py', size=12, font='微软雅黑', indent=True)
para('启动后浏览器会自动打开 http://127.0.0.1:8000（若未自动打开，可手动访问）。', indent=True)
h2('3.4 配置高德 Key')
para('首次使用前，需要到「设置」页填写高德开放平台的 Web 服务 Key（免费申请）。'
     '申请步骤见第七节"常见问题"。', indent=True)
page_break()

# ===================== 四、界面总览 =====================
h1('四、界面总览')
para('软件采用"顶栏 + 导航标签页"的单页布局，共 5 个标签页：路线搜索（首页）、坐标/轨迹、批量导入、今日计划、设置。', indent=True)
img('01_首页_路线搜索.png', '图 1  软件首页（路线搜索）')
para('首页为"路线搜索"页，分为左右两栏：', indent=True)
para('左栏：搜索表单（起点、终点、出行方式、驾车策略、途经点、搜索按钮）与搜索历史列表。', indent=True)
para('右栏：地图（Leaflet + 高德底图瓦片）与结果面板（总距离、预计耗时、导出按钮、分段导航指令表）。', indent=True)
page_break()

# ===================== 五、功能操作说明 =====================
h1('五、功能操作说明')

h2('5.1 路线搜索（首页）')
img('01_首页_路线搜索.png', '图 2  路线搜索界面')
para('操作步骤：', indent=False)
for s in [
    '第 1 步：在"起点""终点"输入框填写地址（如"北京市朝阳区阜通东大街"）或经纬度（如 116.481,39.99，默认 WGS84）。',
    '第 2 步：选择出行方式（驾车 / 步行 / 骑行 / 公交）。',
    '第 3 步（可选）：驾车时可选择驾车策略、添加途经点。',
    '第 4 步：点击"搜索"按钮，地图上即显示路线折线与起终点标记。',
    '第 5 步：结果面板显示总距离、预计耗时与分段导航指令表。',
    '第 6 步：可点击"导出 CSV""导出 GeoJSON"下载路线数据，或"加入今日计划"保存该路线。',
]:
    para(s, size=12, space_after=3)
para('说明：点击起点输入框旁的"📍"按钮，可用浏览器定位一键填入当前位置。', indent=True)

h2('5.2 坐标转换与轨迹录制')
img('02_坐标轨迹.png', '图 3  坐标转换与轨迹录制界面')
para('坐标互转：输入经度、纬度，选择转换方向（WGS84 → GCJ-02 或 GCJ-02 → WGS84），点击"转换"即可得到结果。', indent=True)
para('轨迹录制：点击"开始录制"，应用按固定间隔用浏览器定位采样当前位置；点击"停止"后，可把轨迹导出为 WGS84 的 GeoJSON / CSV。', indent=True)

h2('5.3 批量导入')
img('03_批量导入.png', '图 4  批量导入界面')
para('操作步骤：', indent=False)
for s in [
    '第 1 步：准备 Excel 或 CSV 文件，列顺序为：起点 / 终点 / 出行方式 / 备注。',
    '第 2 步：点击"解析文件"选择文件，下方会列出解析出的起终点。',
    '第 3 步：选择默认出行方式与驾车策略（用于未填方式的行）。',
    '第 4 步：点击"搜索全部路线"，一次搜索所有路线。',
    '第 5 步：可导出里程统计表(CSV)、全部路线(GeoJSON)、全部路线(CSV)。',
]:
    para(s, size=12, space_after=3)

h2('5.4 今日计划')
img('04_今日计划.png', '图 5  今日计划界面')
para('今日计划用于集中管理"当天要跑的路线"：', indent=True)
for s in [
    '加入：在路线搜索页点击"加入今日计划"。',
    '排序：点击每条路线的"↑""↓"调整顺序。',
    '查看：点击"查看"回到首页并显示该路线。',
    '保存：点击"💾 保存"显式保存计划（应用也会在增删改时自动保存）。',
    '导出：一键导出里程统计表 / 全部路线（GeoJSON / CSV）。',
]:
    para(s, size=12, space_after=3)

h2('5.5 设置与备份恢复')
img('05_设置.png', '图 6  设置与备份恢复界面')
para('高德 Key 设置：填写并保存 Web 服务 Key，Key 保存在本地 data/config.json，不会上传。', indent=True)
para('备份 / 恢复：点击"导出备份"把 data/ 与 downloads/ 打包为 zip；点击"导入恢复"选择 zip 恢复数据（恢复前会自动备份当前数据到 backups/）。', indent=True)
page_break()

# ===================== 六、搜索示例 =====================
h1('六、搜索示例（完整操作演示）')
para('下面以"驾车"搜索"北京市朝阳区阜通东大街 → 天安门"为例，演示完整流程：', indent=True)
for s in [
    '第 1 步：启动应用，进入首页。',
    '第 2 步：起点输入"北京市朝阳区阜通东大街"，终点输入"天安门"。',
    '第 3 步：选择出行方式"🚗 驾车"，策略保持默认"最快捷"。',
    '第 4 步：点击"🔍 搜索"。',
    '第 5 步：地图显示蓝色路线（实时路况下按红黄绿着色），右侧显示总距离、预计耗时与分段导航指令表。',
    '第 6 步：点击"导出 GeoJSON"，下载 WGS84 坐标的路线文件；点击"加入今日计划"把路线加入今日计划。',
]:
    para(s, size=12, space_after=3)
img('06_搜索结果.png', '图 7  搜索结果（地图显示路线 + 分段导航表）')
page_break()

# ===================== 七、常见问题 =====================
h1('七、常见问题')
h2('7.1 如何申请高德 Key？')
for s in [
    '① 打开 https://lbs.amap.com/ 注册 / 登录。',
    '② 进入「控制台」，完成个人开发者实名认证（免费）。',
    '③ 「应用管理 → 我的应用 → 创建新应用」。',
    '④ 在应用下「添加 Key」，服务平台务必选「Web 服务」。',
    '⑤ 复制 Key，粘贴到本软件「设置」页保存。',
]:
    para(s, size=12, space_after=3)
h2('7.2 公交/地铁为什么报"没有权限"？')
para('高德的"公交路径规划"接口对个人免费 Key 默认无权限（返回 INSUFFICIENT_PRIVILEGES / 错误码 10012），'
     '需在高德开放平台为该 Key 单独开通该服务（通常需企业认证或单独申请）。驾车 / 步行 / 骑行不受影响。', indent=True)
h2('7.3 步行/骑行距离有限制吗？')
para('有。步行路线规划上限约 100 公里，骑行上限约 500 公里；超出时高德会返回"起终点距离过长"（错误码 20803），'
     '此时请改用驾车。', indent=True)
h2('7.4 输入经纬度按什么坐标系？')
para('按 WGS84 处理。软件内部会自动转换为 GCJ-02 再请求高德，导出时再转换回 WGS84，保证与 GPS / 其它地图工具通用。', indent=True)
h2('7.5 数据保存在哪里？重启会丢吗？')
para('所有数据保存在本机：data/（Key、搜索历史、今日计划）、downloads/（导出文件）、tiles/（瓦片缓存）、backups/（恢复备份）。'
     '刷新、关闭、重启都不会丢失。', indent=True)
page_break()

# ===================== 八、数据与隐私 =====================
h1('八、数据与隐私说明')
para('本软件只在本地运行，不登录、不上传云端，不需要任何账号。', indent=True)
para('仅以下三类操作需要联网调用高德：搜索路线、公交/逆地理编码、加载地图底图（首次）。', indent=True)
para('高德 Key 与所有业务数据均只保存在本机，不会发送到除高德官方接口以外的任何第三方。', indent=True)
para('如需迁移到另一台电脑，可复制 data/ 与 downloads/ 目录，或使用"设置"页的备份/恢复功能。', indent=True)

doc.save(OUT)
print('SAVED', OUT)
