# -*- coding: utf-8 -*-
"""生成结课作业报告（Word，含选题背景/功能/AI使用验证/测试结果/分工）。"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = r'D:\桌面\APP'
SHOTS = os.path.join(ROOT, '交付物', 'screenshots')
OUT = os.path.join(ROOT, '交付物', '结课作业报告.docx')

doc = Document()
for section in doc.sections:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(12)
normal._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')


def zh(run, font='宋体'):
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font)


def para(text, size=12, bold=False, align=None, font='宋体', space_after=6, indent=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    zh(r, font)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.5
    if indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    return p


def h1(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(16)
    r.bold = True
    r.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
    zh(r, '微软雅黑')
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    return p


def h2(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(14)
    r.bold = True
    zh(r, '微软雅黑')
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    return p


def img(path, caption, width_cm=15.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(os.path.join(SHOTS, path), width=Cm(width_cm))
    para(caption, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, font='微软雅黑')


def page_break():
    doc.add_page_break()


# ===================== 封面 =====================
for _ in range(4):
    doc.add_paragraph()
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('结课作业报告'); r.font.size = Pt(30); r.bold = True; zh(r, '微软雅黑')
t2 = doc.add_paragraph(); t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t2.add_run('高德路径搜索与下载工具'); r.font.size = Pt(20); r.bold = True; zh(r, '微软雅黑')
for _ in range(8):
    doc.add_paragraph()
para('姓名：李', size=14, align=WD_ALIGN_PARAGRAPH.CENTER, font='微软雅黑')
para('学号：（请填写）', size=14, align=WD_ALIGN_PARAGRAPH.CENTER, font='微软雅黑')
para('日期：2026 年 10 月', size=14, align=WD_ALIGN_PARAGRAPH.CENTER, font='微软雅黑')
page_break()

# ===================== 目录 =====================
h1('目 录')
for c in ['一、选题背景', '二、功能介绍', '三、系统设计与实现', '四、AI 使用与验证说明',
          '五、测试结果', '六、个人分工', '七、总结与展望', '附录：项目结构与运行方式']:
    para(c, size=13, space_after=6, font='微软雅黑')
page_break()

# ===================== 一、选题背景 =====================
h1('一、选题背景')
h2('1.1 问题来源')
para('在日常出行、外勤走访、配送调度、路线盘点等场景中，人们经常需要查询两点之间的路线，'
     '并了解其距离、耗时以及具体怎么走。市面上的地图应用（如高德地图 APP）虽然可以展示路线，'
     '但往往只能在手机或网页上"看"，难以把路线数据方便地导出、复用或批量处理。', indent=True)
para('此外，路线数据还涉及坐标系问题：高德等国内地图使用 GCJ-02（火星坐标），'
     '而 GPS 设备、开源 GIS 工具普遍使用 WGS84。二者之间存在数百米的偏移，'
     '直接把高德坐标用于 GPS 或 GIS 会定位不准，需要可靠的坐标转换。', indent=True)
h2('1.2 需求分析')
para('综合以上痛点，本项目需要解决以下需求：', indent=True)
for f in [
    '① 能输入地址或经纬度，搜索驾车、步行、骑行、公交等多种出行方式的路线；',
    '② 在地图上直观展示路线，并给出总距离、预计耗时与分段导航指令；',
    '③ 能把路线导出为 CSV、GeoJSON 等通用格式，且坐标统一为 WGS84；',
    '④ 提供 GCJ-02 与 WGS84 的坐标互转工具；',
    '⑤ 支持从 Excel 批量导入多组起终点、一次搜索并统计总里程；',
    '⑥ 数据保存在本机、不依赖云端，刷新与重启后不丢失。',
]:
    para(f, size=12, space_after=3)
h2('1.3 选题意义')
para('本课题通过调用高德开放平台的免费 Web 服务接口，搭建了一个"搜索 + 下载 + 管理"一体化的本地路线工具，'
     '既解决了路线数据导出与二次利用的问题，又实践了 Web 应用开发、接口调用、坐标转换、'
     '数据持久化、自动化测试等多项技术，具有较好的实用价值与学习价值。', indent=True)
page_break()

# ===================== 二、功能介绍 =====================
h1('二、功能介绍')
para('本软件的主要功能模块如下表所示：', indent=True)
# 用表格呈现功能
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr = table.rows[0].cells
hdr[0].text = '功能模块'
hdr[1].text = '说明'
funcs = [
    ('路线搜索', '支持驾车/步行/骑行/公交四种方式，输入地址或经纬度（默认 WGS84），地图显示路线'),
    ('驾车策略', '最快捷、最短距离、少收费、躲避拥堵、不走高速等策略可选'),
    ('途经点', '驾车支持起终点之间插入多个途经点'),
    ('实时路况', '驾车路线按畅通/缓行/拥堵/严重拥堵着色'),
    ('结果展示', '总距离、预计耗时、分段导航指令表'),
    ('导出', '单条路线导出 CSV、GeoJSON（WGS84）'),
    ('坐标互转', 'GCJ-02 ↔ WGS84 互转小工具'),
    ('定位与轨迹', '浏览器定位填起点，轨迹录制并导出 GeoJSON/CSV'),
    ('批量导入', 'Excel/CSV 导入多组起终点，一次搜索多条、里程统计表'),
    ('今日计划', '当日路线任务管理、排序、查看、一键导出'),
    ('备份恢复', '一键打包/恢复本地数据为 zip'),
    ('离线地图', '浏览过的地图瓦片自动缓存，断网可看已缓存区域'),
]
for k, v in funcs:
    row = table.add_row().cells
    row[0].text = k
    row[1].text = v
    for c in row:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.size = Pt(10.5)
                zh(r)
para('', size=6)
img('06_搜索结果.png', '图 1  搜索结果界面')
page_break()

# ===================== 三、系统设计与实现 =====================
h1('三、系统设计与实现')
h2('3.1 总体架构')
para('本软件采用"浏览器前端 + 本地 Python 后端"的轻量级 B/S 架构，整体运行在一台电脑上。', indent=True)
para('后端：使用 Python 标准库 http.server 实现 Web 服务，提供搜索、转换、导入、导出、计划、备份等接口；'
     '通过 urllib 调用高德 Web 服务 API；使用 JSON 文件实现本地持久化。', indent=True)
para('前端：原生 HTML + CSS + JavaScript 单页应用，使用 Leaflet（本地化）与高德底图瓦片展示地图。', indent=True)
h2('3.2 技术选型')
for f in [
    '后端：Python 3.8+（标准库 http.server / urllib / json / csv / io / zipfile），Excel 导入使用 openpyxl。',
    '前端：原生 HTML / CSS / JavaScript + Leaflet 1.9.4（本地化，无 CDN 依赖）。',
    '数据存储：JSON 文件（data/config.json、history.json、plan.json）+ 文件目录（downloads/、tiles/）。',
    '测试：Python 标准库 unittest，运行器 run_tests.py。',
]:
    para(f, size=12, space_after=3)
h2('3.3 关键流程与坐标系处理')
para('一次路线搜索的完整数据流如下：', indent=True)
for f in [
    '① 输入解析：判断输入是"经纬度"还是"地址"；经纬度按 WGS84 解析，地址调用高德地理编码转坐标。',
    '② 坐标转换：WGS84 → GCJ-02，因为高德接口内部使用 GCJ-02。',
    '③ 路线规划：调用高德路径规划接口（驾车 v3 / 步行 v3 / 骑行 v4 / 公交 v3）。',
    '④ 结果处理：解析总距离、耗时、分段指令与坐标点串；同时保留 GCJ-02（地图显示）与 WGS84（导出）两套坐标。',
    '⑤ 展示与导出：地图按 GCJ-02 绘制以贴合高德底图；导出 CSV/GeoJSON 统一为 WGS84。',
]:
    para(f, size=12, space_after=3)
h2('3.4 项目结构')
for f in [
    'app.py：后端主程序与全部接口。',
    'amap_client.py：高德接口封装（地理/逆地理编码、驾车/步行/骑行/公交）。',
    'coord.py：GCJ-02 ↔ WGS84 坐标互转与坐标字符串解析。',
    'storage.py：本地 JSON 持久化。',
    'static/：前端页面（index.html / style.css / app.js / leaflet/）。',
    'tests/：unittest 测试（9 个测试模块）。',
    'run_tests.py：测试运行器。',
]:
    para(f, size=12, space_after=3)
page_break()

# ===================== 四、AI 使用与验证说明 =====================
h1('四、AI 使用与验证说明')
h2('4.1 使用的 AI 工具')
para('本项目开发过程中使用了 DeepSeek 人工智能编程助手，用于辅助需求分析、接口调研、代码编写、测试设计与问题排查。', indent=True)
h2('4.2 AI 辅助的具体内容')
for f in [
    '① 需求澄清与产品设计：协助把零散需求整理为产品需求文档（PRD）与开发计划（DEVELOPMENT_PLAN）。',
    '② 高德接口调研：协助确认驾车/步行/骑行/公交接口的 URL、参数与返回结构。',
    '③ 代码生成：协助编写后端接口、前端页面、坐标转换算法、导出与备份等模块。',
    '④ 测试编写：协助编写 unittest 单元测试与集成测试。',
    '⑤ 问题排查：协助定位"请求失败""pts is not defined""公交无权限"等运行问题。',
]:
    para(f, size=12, space_after=3)
h2('4.3 人工验证与 AI 结果核对')
para('AI 生成的代码与结论并非直接采用，而是经过了以下人工验证：', indent=True)
for f in [
    '① 真实接口联调：用真实高德 Key 逐一验证地理编码、驾车/步行/骑行/公交接口的返回。',
    '② 坐标转换核对：用已知参考值核对 WGS84↔GCJ-02 转换结果（如 116.397428,39.90923 → 116.403672,39.910634）。',
    '③ 自动化测试：编写并运行 78 个单元/集成测试，全部通过。',
    '④ 边界用例：验证未填 Key、起终点相同、距离过长、搜不到路、Key 无权限等异常场景的中文提示。',
    '⑤ 安全校验：下载的官方软件包通过 Authenticode 数字签名与文件大小双重校验。',
]:
    para(f, size=12, space_after=3)
h2('4.4 使用体会')
para('AI 在本项目中主要承担了"快速产出与查错"的角色，显著提高了开发效率；但接口细节、'
     '坐标系约定、异常边界等关键结论仍需人工以真实数据和测试加以确认。'
     '这说明 AI 是高效的工具，但不能替代对结果负责的人工验证。', indent=True)
page_break()

# ===================== 五、测试结果 =====================
h1('五、测试结果')
h2('5.1 测试方案')
para('采用 Python 标准库 unittest 编写自动化测试，运行方式为"python run_tests.py"；'
     '同时对前端 JS 做 node 语法校验，并保留了真实高德接口的联调测试。', indent=True)
h2('5.2 测试覆盖范围')
for f in [
    '坐标互转（GCJ-02 ↔ WGS84）与坐标字符串解析；',
    '驾车/步行/骑行/公交接口返回解析（含路况 tmcs、途经点、错误码翻译）；',
    '本地 JSON 持久化（Key / 历史 / 计划）；',
    '导出（CSV / GeoJSON / 里程统计 / 轨迹）；',
    '备份与恢复（zip 打包、恢复、恢复前备份）；',
    '瓦片缓存（缓存命中、参数校验）；',
    'HTTP 接口（搜索、转换、导入、导出、计划、备份、恢复、瓦片）；',
    '前端资源一致性（CSS 语法、HTML 引用、关键类名、保存按钮绑定）；',
    '真实高德接口联调（地理编码、驾车/步行/骑行）。',
]:
    para(f, size=12, space_after=3)
h2('5.3 测试结果')
para('测试结果统计如下：', indent=True)
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr = table.rows[0].cells
hdr[0].text = '项目'
hdr[1].text = '数值'
hdr[2].text = '结果'
rows = [
    ('测试用例总数', '78', '全部通过'),
    ('测试模块数', '9', '—'),
    ('真实接口联调用例', '5', '通过'),
    ('前端 JS 语法校验', '1 项', '通过'),
]
for a, b, c in rows:
    row = table.add_row().cells
    row[0].text = a; row[1].text = b; row[2].text = c
    for cell in row:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(10.5); zh(r)
para('', size=6)
para('最终全量运行结果为：Ran 78 tests ... OK，退出码 0。', indent=True)
h2('5.4 测试过程中的典型问题')
for f in [
    '① 沙箱环境限制：tempfile.mkdtemp 创建的目录在沙箱下无法再建子目录，改用 os.makedirs 生成临时目录。',
    '② 作用域错误：前端 drawRoute 中变量 pts 被移入 if/else 分支导致未定义，修复为函数级变量 boundsPts。',
    '③ 坐标往返精度断言过严：GCJ-02→WGS84 为近似算法，把断言精度从 6 位放宽到 5 位。',
]:
    para(f, size=12, space_after=3)
page_break()

# ===================== 六、个人分工 =====================
h1('六、个人分工')
para('本项目由个人独立完成，承担了从需求到交付的全部工作，具体分工如下：', indent=True)
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr = table.rows[0].cells
hdr[0].text = '工作阶段'
hdr[1].text = '承担内容'
works = [
    ('需求分析', '梳理功能需求，编写产品需求文档（PRD）与开发计划'),
    ('系统设计', '设计前后端架构、坐标系处理流程、数据存储方案'),
    ('后端开发', '实现 http.server 服务、高德接口封装、坐标转换、导出、备份恢复、瓦片缓存等'),
    ('前端开发', '实现页面布局、地图展示、各功能交互、界面美化'),
    ('测试', '编写并运行 78 个单元/集成测试，处理测试中发现的问题'),
    ('文档与交付', '编写 README、软件说明书、结课报告，打包源码与制作演示材料'),
]
for k, v in works:
    row = table.add_row().cells
    row[0].text = k
    row[1].text = v
    for c in row:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.size = Pt(10.5); zh(r)
para('', size=6)
para('说明：本作业为个人独立完成，因此不存在多人协作的任务拆分；上表按"工作阶段"呈现本人承担的全部工作内容。', indent=True)
page_break()

# ===================== 七、总结与展望 =====================
h1('七、总结与展望')
h2('7.1 总结')
para('本项目完成了一个功能完整、可本地运行的"高德路径搜索与下载工具"，实现了路线搜索、地图展示、'
     '分段导航、CSV/GeoJSON 导出（WGS84）、坐标互转、批量导入与里程统计、今日计划、'
     '轨迹录制、备份恢复、离线瓦片缓存等功能，并通过 78 个自动化测试保证了核心逻辑的正确性。', indent=True)
para('在开发过程中，实践了 Web 应用开发、第三方接口调用、坐标系转换、数据持久化、'
     '自动化测试、以及"AI 辅助 + 人工验证"的协作方式，收获较大。', indent=True)
h2('7.2 不足与展望')
for f in [
    '公交/地铁功能受限于高德对个人免费 Key 的权限，需企业认证后才能真正使用。',
    '离线地图目前为"边浏览边缓存"的透明缓存，可进一步支持预先按区域下载。',
    '前端缺少单元级自动化测试（当前为资源一致性与语法校验），可引入 Playwright 等做端到端测试。',
    '后续可考虑登录、云同步、多设备等能力（与"本地运行"的定位需进一步权衡）。',
]:
    para(f, size=12, space_after=3)
page_break()

# ===================== 附录 =====================
h1('附录：项目结构与运行方式')
h2('A.1 项目结构')
for f in [
    'app.py / amap_client.py / coord.py / storage.py —— 后端核心模块',
    'static/ —— 前端页面（index.html、style.css、app.js、leaflet/）',
    'tests/ —— 9 个 unittest 测试模块',
    'run_tests.py —— 测试运行器',
    'requirements.txt —— 依赖清单（openpyxl）',
    'README.md / PRD.md / DEVELOPMENT_PLAN.md —— 文档',
]:
    para(f, size=12, space_after=3)
h2('A.2 运行方式')
para('pip install -r requirements.txt', size=12, font='微软雅黑', indent=True)
para('python app.py', size=12, font='微软雅黑', indent=True)
para('浏览器打开 http://127.0.0.1:8000，在「设置」页填写高德 Web 服务 Key 后即可使用。', indent=True)
h2('A.3 测试方式')
para('python run_tests.py', size=12, font='微软雅黑', indent=True)
para('输出 "Ran 78 tests ... OK" 即为通过。', indent=True)

doc.save(OUT)
print('SAVED', OUT)
